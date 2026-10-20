from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

from inventory_launcher import format_count_text, read_fixed_parameters, read_inventory_entries

PREVIEW_ROWS = 10


def resolve_sheet_path() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    return Path("inventario_13.xlsx")


def print_raw_cells(sheet_path: Path, row_limit: int) -> None:
    workbook = openpyxl.load_workbook(filename=sheet_path, read_only=True)
    worksheet = workbook.active
    if worksheet is None:
        print("Planilha sem aba ativa")
        return None
    last_row = worksheet.max_row or 0
    for row_index in range(1, min(last_row, row_limit) + 1):
        code_value = repr(worksheet.cell(row=row_index, column=1).value)
        count_value = repr(worksheet.cell(row=row_index, column=2).value)
        print(f"linha {row_index}: A={code_value} B={count_value}")
    workbook.close()
    return None


def main() -> int:
    sheet_path = resolve_sheet_path()
    if not sheet_path.exists():
        print(f"Arquivo não encontrado: {sheet_path.resolve()}")
        return 1
    print_raw_cells(sheet_path, PREVIEW_ROWS)
    try:
        branch_code, warehouse_code = read_fixed_parameters(sheet_path)
        print(f"C2/D2: filial={branch_code} deposito={warehouse_code}")
    except ValueError as error:
        print(f"C2/D2: {error}")
    try:
        entries = read_inventory_entries(sheet_path, 2)
        print(f"Leitura OK: {len(entries)} itens")
        for entry in entries[:PREVIEW_ROWS]:
            print(f"{entry.item_code}\t{format_count_text(entry.count_quantity)}")
    except ValueError as error:
        print(f"Erro de leitura: {error}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())