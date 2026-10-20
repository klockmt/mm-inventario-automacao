from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

from inventory_launcher import (
    CODE_COLUMN,
    COUNT_COLUMN,
    ENTER_KEY,
    AutomationStoppedError,
    InventoryEntry,
    LaunchConfig,
    LaunchResult,
    PyAutoGuiWriter,
    build_inventory_entry,
    format_item_code,
    format_launch_result,
    is_blank_value,
    launch_entry,
    launch_inventory_entries,
    parse_config,
    parse_count,
    print_inventory_entries,
    read_inventory_entries,
    validate_start_row,
    REVERT_COMBO_KEYS,
    BRANCH_COLUMN,
    WAREHOUSE_COLUMN,
    format_revert_result,
    read_fixed_parameters,
    revert_entry,
    revert_inventory_entries,
    format_count_text,
)

HEADER_ROW = 1
FIRST_ROW = 2
ITEM_CODE = "12345"
COUNT_QUANTITY = 7


class FakeKeyWriter:
    def __init__(self) -> None:
        self.actions: list[str] = []
        self.paused_seconds: list[float] = []

    def type_text(self, text: str) -> None:
        self.actions.append(f"type:{text}")
        return None

    def press_key(self, key: str) -> None:
        self.actions.append(f"press:{key}")
        return None

    def press_combo(self, keys: Sequence[str]) -> None:
        self.actions.append("combo:" + "+".join(keys))
        return None

    def pause_seconds(self, seconds: float) -> None:
        self.paused_seconds.append(seconds)
        return None


class FakeStoppingKeyWriter:
    def __init__(self, stop_text: str) -> None:
        self.actions: list[str] = []
        self.paused_seconds: list[float] = []
        self.stop_text = stop_text
        return None

    def type_text(self, text: str) -> None:
        if text == self.stop_text:
            raise AutomationStoppedError("parada")
        self.actions.append(f"type:{text}")
        return None

    def press_key(self, key: str) -> None:
        self.actions.append(f"press:{key}")
        return None

    def press_combo(self, keys: Sequence[str]) -> None:
        self.actions.append("combo:" + "+".join(keys))
        return None

    def pause_seconds(self, seconds: float) -> None:
        self.paused_seconds.append(seconds)
        return None


class FakeGuiKeyboard:
    def __init__(self) -> None:
        self.typed_messages: list[tuple[str, float]] = []
        self.pressed_keys: list[str] = []
        self.pressed_combos: list[tuple[str, ...]] = []

    def typewrite(self, message: str, interval: float) -> None:
        self.typed_messages.append((message, interval))
        return None

    def press(self, key: str) -> None:
        self.pressed_keys.append(key)
        return None

    def hotkey(self, *keys: str) -> None:
        self.pressed_combos.append(keys)
        return None


class FakeSleep:
    def __init__(self) -> None:
        self.calls: list[float] = []

    def __call__(self, seconds: float) -> None:
        self.calls.append(seconds)
        return None


def save_inventory_sheet(
    sheet_path: Path,
    rows: list[tuple[object, object]],
    branch_code: object = None,
    warehouse_code: object = None,
) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.cell(row=HEADER_ROW, column=CODE_COLUMN, value="codigo")
    worksheet.cell(row=HEADER_ROW, column=COUNT_COLUMN, value="contagem")
    if branch_code is not None:
        worksheet.cell(row=FIRST_ROW, column=BRANCH_COLUMN, value=branch_code)
    if warehouse_code is not None:
        worksheet.cell(row=FIRST_ROW, column=WAREHOUSE_COLUMN, value=warehouse_code)
    for offset, (code_value, count_value) in enumerate(rows):
        worksheet.cell(row=FIRST_ROW + offset, column=CODE_COLUMN, value=code_value)
        worksheet.cell(row=FIRST_ROW + offset, column=COUNT_COLUMN, value=count_value)
    workbook.save(sheet_path)
    return None


def expected_entry_actions(item_code: str, count_text: str) -> list[str]:
    enter_action = f"press:{ENTER_KEY}"
    return [
        enter_action,
        f"type:{item_code}",
        enter_action,
        enter_action,
        enter_action,
        enter_action,
        f"type:{count_text}",
        enter_action,
        f"type:{count_text}",
        enter_action,
        enter_action,
    ]


@pytest.fixture()
def inventory_sheet_path(tmp_path: Path) -> Path:
    return tmp_path / "inventario.xlsx"


def test_validate_start_row_rejects_header_row() -> None:
    with pytest.raises(ValueError):
        validate_start_row(1)
    return None


def test_is_blank_value_identifies_none_and_blank_text() -> None:
    assert is_blank_value(None) is True
    assert is_blank_value("   ") is True
    assert is_blank_value("0") is False
    assert is_blank_value(0) is False
    return None


def test_format_item_code_keeps_text_and_normalizes_numbers() -> None:
    assert format_item_code(" 00042 ") == "00042"
    assert format_item_code(42) == "42"
    assert format_item_code(42.0) == "42"
    assert format_item_code(None) == ""
    return None


def test_parse_count_accepts_integer_and_text() -> None:
    assert parse_count(7) == 7.0
    assert parse_count(7.0) == 7.0
    assert parse_count(7.5) == 7.5
    assert parse_count("8") == 8.0
    assert parse_count("1,5") == 1.5
    assert parse_count(None) == 0.0
    return None


@pytest.mark.parametrize("invalid_count", [True, "abc", "-1"])
def test_parse_count_rejects_invalid_values(invalid_count: object) -> None:
    with pytest.raises(ValueError):
        parse_count(invalid_count)
    return None


def test_format_count_text_keeps_integers_and_uses_comma_for_fractions() -> None:
    assert format_count_text(7.0) == "7"
    assert format_count_text(1.5) == "1,5"
    assert format_count_text(832.0688566800001) == "832,0689"
    return None


def test_launch_entry_types_fractional_count_with_comma() -> None:
    writer = FakeKeyWriter()
    launch_entry(writer, InventoryEntry(ITEM_CODE, 1.5))
    assert writer.actions == expected_entry_actions(ITEM_CODE, "1,5")
    return None


def test_build_inventory_entry_raises_when_count_missing() -> None:
    with pytest.raises(ValueError):
        build_inventory_entry(FIRST_ROW, ITEM_CODE, None)
    return None


def test_build_inventory_entry_accepts_zero_count() -> None:
    entry = build_inventory_entry(FIRST_ROW, ITEM_CODE, 0)
    assert entry == InventoryEntry(ITEM_CODE, 0)
    return None


def test_read_inventory_entries_reads_rows_after_header(inventory_sheet_path: Path) -> None:
    save_inventory_sheet(inventory_sheet_path, [(ITEM_CODE, COUNT_QUANTITY), ("00042", 3)])
    entries = read_inventory_entries(inventory_sheet_path, FIRST_ROW)
    assert entries == [
        InventoryEntry(ITEM_CODE, COUNT_QUANTITY),
        InventoryEntry("00042", 3),
    ]
    return None


def test_read_inventory_entries_skips_blank_rows(inventory_sheet_path: Path) -> None:
    save_inventory_sheet(
        inventory_sheet_path,
        [(ITEM_CODE, COUNT_QUANTITY), (None, None), ("A", 1)],
    )
    entries = read_inventory_entries(inventory_sheet_path, FIRST_ROW)
    assert entries == [
        InventoryEntry(ITEM_CODE, COUNT_QUANTITY),
        InventoryEntry("A", 1),
    ]
    return None


def test_read_inventory_entries_raises_when_code_missing(inventory_sheet_path: Path) -> None:
    save_inventory_sheet(inventory_sheet_path, [(None, COUNT_QUANTITY)])
    with pytest.raises(ValueError):
        read_inventory_entries(inventory_sheet_path, FIRST_ROW)
    return None


def test_launch_entry_uses_expected_sequence() -> None:
    writer = FakeKeyWriter()
    launch_entry(writer, InventoryEntry(ITEM_CODE, COUNT_QUANTITY))
    assert writer.actions == expected_entry_actions(ITEM_CODE, str(COUNT_QUANTITY))
    return None


def test_launch_inventory_entries_waits_and_launches_all() -> None:
    writer = FakeKeyWriter()
    entries = [InventoryEntry("A", 1), InventoryEntry("B", 2)]
    result = launch_inventory_entries(writer, entries, 3.0)
    assert result.status == "concluido"
    assert result.launched_entries == tuple(entries)
    assert writer.paused_seconds == [3.0]
    assert writer.actions == expected_entry_actions("A", "1") + expected_entry_actions("B", "2")
    return None


def test_pyautogui_writer_types_and_presses(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_sleep = FakeSleep()
    monkeypatch.setattr("inventory_launcher.sleep", fake_sleep)
    keyboard = FakeGuiKeyboard()
    writer = PyAutoGuiWriter(0.0, keyboard)
    writer.type_text("A")
    writer.press_key(ENTER_KEY)
    assert keyboard.typed_messages == [("A", 0.0)]
    assert keyboard.pressed_keys == [ENTER_KEY]
    return None


def test_pyautogui_writer_ignores_empty_text_and_key(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_sleep = FakeSleep()
    monkeypatch.setattr("inventory_launcher.sleep", fake_sleep)
    keyboard = FakeGuiKeyboard()
    writer = PyAutoGuiWriter(0.0, keyboard)
    writer.type_text("")
    writer.press_key("")
    assert keyboard.typed_messages == []
    assert keyboard.pressed_keys == []
    return None


def test_pyautogui_writer_pause_calls_sleep(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_sleep = FakeSleep()
    monkeypatch.setattr("inventory_launcher.sleep", fake_sleep)
    keyboard = FakeGuiKeyboard()
    writer = PyAutoGuiWriter(0.0, keyboard)
    writer.pause_seconds(1.5)
    assert fake_sleep.calls == [1.5]
    return None


def test_print_inventory_entries_outputs_code_and_count(capsys: pytest.CaptureFixture[str]) -> None:
    print_inventory_entries([InventoryEntry("A", 1), InventoryEntry("B", 2)])
    captured = capsys.readouterr()
    assert captured.out == "A\t1\nB\t2\n"
    return None


def test_parse_config_builds_launch_config(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "inventory_launcher.py",
            "--planilha",
            "inventario.xlsx",
            "--linha-inicial",
            "3",
            "--espera-inicial",
            "5.0",
            "--pausa-tecla",
            "0.1",
            "--simular",
        ],
    )
    config = parse_config()
    assert config == LaunchConfig(
        sheet_path=Path("inventario.xlsx"),
        start_row=3,
        initial_wait_seconds=5.0,
        key_pause_seconds=0.1,
        simulate=True,
    )
    return None

def test_launch_inventory_entries_returns_partial_result_when_stopped() -> None:
    writer = FakeStoppingKeyWriter(stop_text="B")
    entries = [InventoryEntry("A", 1), InventoryEntry("B", 2)]
    result = launch_inventory_entries(writer, entries, 0.0)
    assert result.status == "parado"
    assert result.launched_entries == (InventoryEntry("A", 1),)
    assert writer.paused_seconds == [0.0]
    return None


def test_format_launch_result_shows_completed_summary() -> None:
    result = LaunchResult(
        status="concluido",
        launched_entries=(InventoryEntry("A", 1), InventoryEntry("B", 2)),
    )
    expected_text = "Lançamento concluído.\nLinhas lançadas: 2\nÚltimo item: B | 2"
    assert format_launch_result(result) == expected_text
    return None


def test_format_launch_result_shows_stopped_without_entries() -> None:
    result = LaunchResult(status="parado", launched_entries=())
    assert format_launch_result(result) == "Lançamento parado.\nNenhuma linha lançada."
    return None

def expected_revert_actions(item_code: str, branch_code: str, warehouse_code: str) -> list[str]:
    enter_action = f"press:{ENTER_KEY}"
    return [
        f"type:{branch_code}",
        enter_action,
        f"type:{item_code}",
        enter_action,
        f"type:{warehouse_code}",
        enter_action,
        "combo:shift+f2",
        enter_action,
    ]


def test_revert_entry_uses_expected_sequence() -> None:
    writer = FakeKeyWriter()
    revert_entry(writer, InventoryEntry(ITEM_CODE, COUNT_QUANTITY), "01", "05")
    assert writer.actions == expected_revert_actions(ITEM_CODE, "01", "05")
    return None


def test_revert_inventory_entries_returns_partial_result_when_stopped() -> None:
    writer = FakeStoppingKeyWriter(stop_text="B")
    entries = [InventoryEntry("A", 1), InventoryEntry("B", 2)]
    result = revert_inventory_entries(writer, entries, "01", "05", 0.0)
    assert result.status == "parado"
    assert result.launched_entries == (InventoryEntry("A", 1),)
    return None


def test_read_fixed_parameters_reads_branch_and_warehouse(inventory_sheet_path: Path) -> None:
    save_inventory_sheet(
        inventory_sheet_path,
        [(ITEM_CODE, COUNT_QUANTITY)],
        branch_code="01",
        warehouse_code="05",
    )
    branch_code, warehouse_code = read_fixed_parameters(inventory_sheet_path)
    assert branch_code == "01"
    assert warehouse_code == "05"
    return None


def test_read_fixed_parameters_raises_when_missing(inventory_sheet_path: Path) -> None:
    save_inventory_sheet(inventory_sheet_path, [(ITEM_CODE, COUNT_QUANTITY)])
    with pytest.raises(ValueError):
        read_fixed_parameters(inventory_sheet_path)
    return None


def test_format_revert_result_shows_summary() -> None:
    result = LaunchResult(status="concluido", launched_entries=(InventoryEntry("A", 1),))
    assert format_revert_result(result) == "Reversão concluída.\nLinhas revertidas: 1\nÚltimo item: A | 1"
    return None


def test_pyautogui_writer_presses_combo(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_sleep = FakeSleep()
    monkeypatch.setattr("inventory_launcher.sleep", fake_sleep)
    keyboard = FakeGuiKeyboard()
    writer = PyAutoGuiWriter(0.0, keyboard)
    writer.press_combo(REVERT_COMBO_KEYS)
    assert keyboard.pressed_combos == [("shift", "f2")]
    return None

def test_format_count_text_keeps_integers_and_uses_comma_for_fractions() -> None:
    assert format_count_text(7.0) == "7"
    assert format_count_text(1.5) == "1,5"
    return None


def test_launch_entry_types_fractional_count_with_comma() -> None:
    writer = FakeKeyWriter()
    launch_entry(writer, InventoryEntry(ITEM_CODE, 1.5))
    assert writer.actions == expected_entry_actions(ITEM_CODE, "1,5")
    return None